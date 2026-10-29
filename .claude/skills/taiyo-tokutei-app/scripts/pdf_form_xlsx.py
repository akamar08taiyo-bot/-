# PDF の様式（市の Word の様式を LibreOffice で PDF にしたもの）の1ページ（か、いくつかのページの範囲を縦につないだもの）を、
# 同じ形の Excel のシートにする（北九州市の住宅改修・特定福祉用具の申請書は Word の様式しかないため。2026-10-29）。
#   線（罫線）と文字の位置 → 列と行の境目（線の位置を優先し、近い境目はまとめる）→ セルの罫線（太さ・点線）、
#   文字 → 文字の位置にぴったりの結合セル（左ぞろえ・上下中央。字体・大きさ・太字は原本のまま）、
#   縦書き → 囲まれた欄を結合して縦書き、塗り → セルの塗り、文字のない囲まれた欄 → 結合（アプリが書く欄）。
# spec（JSON）：
#   {"pdf": "…pdf", "out": "…xlsx", "map": "…json",
#    "sheets": [{"name": "シート名", "segs": [[ページ(0から), 上pt|null, 下pt|null], …],
#                "font": "ＭＳ ゴシック", "fonts": {"IPAPGothic": "HG丸ｺﾞｼｯｸM-PRO"}, "bold": ["太字の文字", …], "boldFonts": ["太字を描いた PDF の字体"], "landscape": false,
#                "drop": [[x0, 上, x1, 下], …],   ← 入れない範囲（記入例の吹き出しなど）
#                "fields": [{"name": "name", "at": [x, y] か "box": [x0, 上, x1, 下], "x0"/"x1"/"top"/"bottom": 上書き,
#                            "h": "left|center|right", "v": "center|top|bottom", "size": 10.5, "wrap": true, "indent": 1}]}]}
#   map には、欄の名前 → セルの範囲（"C5:F7"）と、文字 → セルの番地（確かめるため）を書く。
# usage: python3 pdf_form_xlsx.py spec.json
import json
import sys
from collections import Counter

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

LINE_TOL = 1.2  # 線の位置をまとめる幅（pt）
MIN_GAP = 0.9  # 境目どうしの最小の間（pt。これより近い文字の境目は線の境目に寄せる）


def cluster(vals, tol):
    vals = sorted(vals)
    out = []
    group = []
    for v in vals:
        if group and v - group[-1] > tol:
            out.append(sum(group) / len(group))
            group = []
        group.append(v)
    if group:
        out.append(sum(group) / len(group))
    return out


def nearest(bs, v):
    lo, hi = 0, len(bs) - 1
    while lo < hi:
        mid = (lo + hi) // 2
        if bs[mid] < v:
            lo = mid + 1
        else:
            hi = mid
    if lo > 0 and abs(bs[lo - 1] - v) <= abs(bs[lo] - v):
        return lo - 1
    return lo


def white(c):
    if c is None:
        return True
    if isinstance(c, (int, float)):
        return c >= 0.98
    if len(c) == 1:
        return c[0] >= 0.98
    if len(c) == 3:
        return min(c) >= 0.98
    if len(c) == 4:  # CMYK
        return max(c) <= 0.02
    return True


def hexcolor(c):
    if isinstance(c, (int, float)):
        c = (c, c, c)
    if len(c) == 1:
        c = (c[0], c[0], c[0])
    if len(c) == 4:
        k = c[3]
        c = tuple((1 - x) * (1 - k) for x in c[:3])
    return ''.join(f'{round(max(0, min(1, x)) * 255):02X}' for x in c)


def primitives(pdf, segs, drop, join=None, clip=None):
    """線（横 H・縦 V）・文字・塗りを、ページの範囲を縦につないだ座標で集める。
    つないだ所（前の範囲の下・次の範囲の上）の横の線は、ページの切れ目でも本当の行の境目のことが多いので残し（2つの線は同じ高さになる）、
    join の横の範囲（[[x0, x1], …]。ページをまたいで続く欄）だけ外す"""
    H, V, chars, fills = [], [], [], []
    yoff = 0.0
    width = 0.0
    for si, sg in enumerate(segs):
        # 4つめ（あれば）：この範囲の前にあける高さ（pt）。LibreOffice で重なった所（委任状の枠と注意書き）を Word のように離す
        pg, top, bottom = sg[0], sg[1], sg[2]
        if len(sg) > 3:
            yoff += sg[3]
        p = pdf.pages[pg]
        width = max(width, p.width)
        top = 0.0 if top is None else top
        bottom = p.height if bottom is None else bottom
        dy = yoff - top

        def dropped(x0, t, x1, b):
            return any(x0 >= d[0] - 0.5 and x1 <= d[2] + 0.5 and t >= d[1] - 0.5 and b <= d[3] + 0.5 for d in drop if len(d) == 4) or any(
                d[4] == pg and x0 >= d[0] - 0.5 and x1 <= d[2] + 0.5 and t >= d[1] - 0.5 and b <= d[3] + 0.5 for d in drop if len(d) == 5)

        def seg(x0, y0, x1, y1, lw, dash):
            if abs(y1 - y0) < 0.6:
                y = (y0 + y1) / 2
                cut = (si > 0 and abs(y - top) < 0.8) or (si < len(segs) - 1 and abs(y - bottom) < 0.8)
                if top - 0.3 <= y <= bottom + 0.3 and abs(x1 - x0) > 0.8 and not dropped(min(x0, x1), y, max(x0, x1), y):
                    parts = [(min(x0, x1), max(x0, x1))]
                    for ja, jb in (join or []) if cut else []:
                        parts = [q for a, b in parts for q in ((a, min(b, ja)), (max(a, jb), b)) if q[1] - q[0] > 0.8]
                    for c in clip or []:
                        if (len(c) < 5 or c[4] == pg) and c[1] <= y <= c[3]:
                            parts = [q for a, b in parts for q in ((a, min(b, c[0])), (max(a, c[2]), b)) if q[1] - q[0] > 0.8]
                    for a, b in parts:
                        H.append([a, b, y + dy, lw, dash])
            elif abs(x1 - x0) < 0.6:
                a, b = sorted((y0, y1))
                a, b = max(a, top), min(b, bottom)
                parts = [(a, b)]
                # clip：その範囲の線だけ消す（LibreOffice で文字がはみ出して、枠がほかの欄に食いこんだ所）
                for c in clip or []:
                    if (len(c) < 5 or c[4] == pg) and c[0] <= x0 <= c[2]:
                        parts = [q for u, w in parts for q in ((u, min(w, c[1])), (max(u, c[3]), w)) if q[1] - q[0] > 0.8]
                for a, b in parts:
                    if b - a > 0.8 and not dropped(x0, a, x0, b):
                        V.append([(x0 + x1) / 2, a + dy, b + dy, lw, dash])

        for l in p.lines:
            pts = l['pts']
            for (ax, ay), (bx, by) in zip(pts, pts[1:]):
                seg(ax, ay, bx, by, l['linewidth'], l.get('dash'))
        for c in p.curves:
            if not c.get('stroke', True):
                continue
            pts = c['pts']
            xs, ys = [q[0] for q in pts], [q[1] for q in pts]
            x0, x1, t, b = min(xs), max(xs), min(ys), max(ys)
            closed = abs(pts[0][0] - pts[-1][0]) < 0.5 and abs(pts[0][1] - pts[-1][1]) < 0.5
            # 角の丸い四角（概要書の裏の見出しの枠）：Excel のセルの罫線では丸くできないので、まわりの四角にする
            # （点をつなぐと、上下の線と、角の丸みのあいだの短い縦の線だけになる）
            if closed and x1 - x0 > 20 and b - t > 10 and all(min(q[0] - x0, x1 - q[0], q[1] - t, b - q[1]) < 12 for q in pts):
                for a in ((x0, t, x1, t), (x0, b, x1, b), (x0, t, x0, b), (x1, t, x1, b)):
                    seg(*a, c['linewidth'], c.get('dash'))
                continue
            for (ax, ay), (bx, by) in zip(pts, pts[1:]):
                seg(ax, ay, bx, by, c['linewidth'], c.get('dash'))
        for r in p.rects:
            x0, x1, t, b = r['x0'], r['x1'], r['top'], r['bottom']
            w, h = x1 - x0, b - t
            if r.get('fill') and not white(r.get('non_stroking_color')):
                if h < 2.5 and w > 3:  # 線の代わりの細い四角
                    seg(x0, (t + b) / 2, x1, (t + b) / 2, h, None)
                    continue
                if w < 2.5 and h > 3:
                    seg((x0 + x1) / 2, t, (x0 + x1) / 2, b, w, None)
                    continue
                tt, bb = max(t, top), min(b, bottom)
                if bb > tt and not dropped(x0, tt, x1, bb):
                    fills.append([x0, tt + dy, x1, bb + dy, hexcolor(r['non_stroking_color'])])
            if r.get('stroke'):
                lw, dash = r['linewidth'], r.get('dash')
                seg(x0, t, x1, t, lw, dash)
                seg(x0, b, x1, b, lw, dash)
                seg(x0, t, x0, b, lw, dash)
                seg(x1, t, x1, b, lw, dash)
        for ch in p.chars:
            if not ch['text'].strip():
                continue
            cy = (ch['top'] + ch['bottom']) / 2
            if not (top <= cy <= bottom) or dropped(ch['x0'], ch['top'], ch['x1'], ch['bottom']):
                continue
            chars.append({'t': ch['text'], 'x0': ch['x0'], 'x1': ch['x1'], 'top': ch['top'] + dy, 'bottom': ch['bottom'] + dy,
                          's': round(ch['size'] * 2) / 2, 'f': ch['fontname'].split('+')[-1]})
        yoff += bottom - top
    return H, V, chars, fills, width, yoff


def v_between(V, xa, xb, y):
    return any(xa - 0.3 <= v[0] <= xb + 0.3 and v[1] - 0.5 <= y <= v[2] + 0.5 for v in V)


def h_between(H, ya, yb, x):
    return any(ya - 0.3 <= h[2] <= yb + 0.3 and h[0] - 0.5 <= x <= h[1] + 0.5 for h in H)


def text_items(chars, H, V, bold, bold_fonts=()):
    """文字を、横書きのかたまり（1行）と縦書きのかたまりにする"""
    for c in chars:
        c['cx'] = (c['x0'] + c['x1']) / 2
        c['cy'] = (c['top'] + c['bottom']) / 2
    # 横に並ぶ字がない字を縦書きの候補に
    def has_h_neighbor(c):
        for d in chars:
            if d is c or abs(d['cy'] - c['cy']) > 0.3 * max(c['s'], d['s']):
                continue
            gap = d['x0'] - c['x1'] if d['x0'] >= c['x0'] else c['x0'] - d['x1']
            # 「令和　　年　　月　　日」のように字の間をあけた行の字も横書き（4字分まで）
            if -0.5 * max(c['s'], d['s']) <= gap <= 4 * max(c['s'], d['s']):
                # あいだに縦の線があれば別の欄（字の箱が線にかかることがあるので、字のまん中どうしのあいだで見る）
                xa, xb = sorted((c['cx'], d['cx']))
                if not v_between(V, xa, xb, c['cy']):
                    return True
        return False

    lone = [c for c in chars if not has_h_neighbor(c)]
    lone.sort(key=lambda c: (c['top'], c['x0']))
    used = set()
    verticals = []
    # 縦書きの字は、LibreOffice の PDF では字ごとに左右に字の幅ほどずれることがある（「前」「の」だけ左）→ 横の範囲が重なるか接していれば同じ列
    for c in lone:
        if id(c) in used:
            continue
        run = [c]
        lo, hi = c['x0'], c['x1']
        for d in lone:
            if id(d) in used or d is c or d['top'] < run[-1]['bottom'] - 0.5 or d['s'] != c['s']:
                continue
            if not (d['x0'] <= hi + 0.2 * c['s'] and d['x1'] >= lo - 0.2 * c['s']):
                continue
            last = run[-1]
            if d['top'] - last['bottom'] > 6 * c['s'] or h_between(H, last['bottom'], d['top'], (max(lo, d['x0']) + min(hi, d['x1'])) / 2):
                break
            run.append(d)
            lo, hi = min(lo, d['x0']), max(hi, d['x1'])
        gaps = [b['top'] - a['bottom'] for a, b in zip(run, run[1:])]
        # 縦書き：字が詰まって並ぶか、3字以上が同じ間で並ぶ（均等割付）。離れた2字・間のそろわない字は、べつべつの横書き
        tight = gaps and max(gaps) <= 1.2 * c['s']
        even = len(gaps) >= 2 and max(gaps) - min(gaps) <= 0.25 * max(max(gaps), c['s'])
        same = len({d['t'] for d in run}) == 1  # 「日」「日」のように同じ字だけ（折り返した行の字が縦に並んだもの）は縦書きではない
        if len(run) >= 2 and (tight or even) and not same:
            for d in run:
                used.add(id(d))
            verticals.append(run)
    rest = [c for c in chars if id(c) not in used]
    # 横書きの行
    rest.sort(key=lambda c: (c['cy'], c['x0']))
    lines = []
    for c in rest:
        for ln in lines:
            if abs(ln['cy'] - c['cy']) <= 0.25 * max(ln['s'], c['s']) and ln['s'] == c['s']:
                ln['chars'].append(c)
                break
        else:
            lines.append({'cy': c['cy'], 's': c['s'], 'chars': [c]})
    items = []
    bold_set = [b.replace('　', '').replace(' ', '') for b in bold]
    for ln in lines:
        cs = sorted(ln['chars'], key=lambda c: c['x0'])
        frag = [cs[0]]
        frags = [frag]
        for a, b in zip(cs, cs[1:]):
            s = max(a['s'], b['s'])
            gap = b['x0'] - a['x1']
            if gap > 1.5 * s or v_between(V, a['cx'], b['cx'], (a['cy'] + b['cy']) / 2):
                frag = [b]
                frags.append(frag)
            else:
                frag.append(b)
        for fr in frags:
            text = fr[0]['t']
            for a, b in zip(fr, fr[1:]):
                gap = b['x0'] - a['x1']
                k = round(gap / (0.5 * b['s'])) if gap > 0.2 * b['s'] else 0
                text += '　' * (k // 2) + ' ' * (k % 2) + b['t']
            plain = text.replace('　', '').replace(' ', '')
            items.append({'kind': 'h', 'text': text, 'x0': fr[0]['x0'], 'x1': fr[-1]['x1'], 'top': min(c['top'] for c in fr),
                          'bottom': max(c['bottom'] for c in fr), 's': Counter(c['s'] for c in fr).most_common(1)[0][0],
                          'f': Counter(c['f'] for c in fr).most_common(1)[0][0],
                          'b': any(plain == bs or (len(plain) >= 5 and plain in bs) for bs in bold_set) or
                          Counter(c['f'] for c in fr).most_common(1)[0][0] in bold_fonts})
    for run in verticals:
        # 縦書きの字の箱は左に字の幅までずれる → 列の位置は右の端から半字
        items.append({'kind': 'v', 'text': ''.join(c['t'] for c in run), 'vx': max(c['x1'] for c in run) - 0.5 * run[0]['s'],
                      'x0': min(c['x0'] for c in run), 'x1': max(c['x1'] for c in run),
                      'top': run[0]['top'], 'bottom': run[-1]['bottom'], 's': Counter(c['s'] for c in run).most_common(1)[0][0],
                      'f': Counter(c['f'] for c in run).most_common(1)[0][0], 'b': False})
    for it in items:
        it['cy'] = (it['top'] + it['bottom']) / 2
    return items


def regions(XB, YB, hb, vb):
    """線で囲まれた欄（グリッドのセルを、あいだに線がないものどうしつなぐ）→ 長方形のもの（i0, j0, i1, j1。i1・j1 は含まない）"""
    nx, ny = len(XB) - 1, len(YB) - 1
    parent = list(range(nx * ny))

    def find(a):
        while parent[a] != a:
            parent[a] = parent[parent[a]]
            a = parent[a]
        return a

    def union(a, b):
        ra, rb = find(a), find(b)
        if ra != rb:
            parent[ra] = rb

    for j in range(ny):
        for i in range(nx):
            if i + 1 < nx and (i + 1, j) not in vb:
                union(j * nx + i, j * nx + i + 1)
            if j + 1 < ny and (j + 1, i) not in hb:
                union(j * nx + i, (j + 1) * nx + i)
    comps = {}
    for j in range(ny):
        for i in range(nx):
            comps.setdefault(find(j * nx + i), []).append((i, j))
    out = []
    for cells in comps.values():
        i0 = min(c[0] for c in cells)
        i1 = max(c[0] for c in cells) + 1
        j0 = min(c[1] for c in cells)
        j1 = max(c[1] for c in cells) + 1
        if len(cells) != (i1 - i0) * (j1 - j0):
            continue
        closed = all((i0, j) in vb and (i1, j) in vb for j in range(j0, j1)) and all((j0, i) in hb and (j1, i) in hb for i in range(i0, i1))
        if closed:
            out.append((i0, j0, i1, j1))
    return out


def side_of(lw, dash):
    if dash and dash[0] and any(d > 0 for d in dash[0]):
        on = dash[0][0]
        if on <= 1.2:
            return 'dotted' if lw < 1.15 else 'mediumDashDotDot'
        return 'dashed' if lw < 1.15 else 'mediumDashed'
    if lw >= 2.0:
        return 'thick'
    if lw >= 1.15:
        return 'medium'
    if lw <= 0.3:
        return 'hair'
    return 'thin'


WEIGHT = {'hair': 0, 'dotted': 1, 'dashed': 2, 'thin': 3, 'mediumDashDotDot': 4, 'mediumDashed': 5, 'medium': 6, 'thick': 7}


def build_sheet(ws, pdf, sh):
    H, V, chars, fills, width, height = primitives(pdf, [tuple(s) for s in sh['segs']], sh.get('drop', []), sh.get('join'), sh.get('clip'))
    # 線の端を、交わる線のまん中に寄せる（太い線の端は線の外側で終わるので、1pt ほどの細い行・列ができる）
    for v in V:
        for k in (1, 2):
            near = [h for h in H if h[0] - 1 <= v[0] <= h[1] + 1 and abs(h[2] - v[k]) <= max(h[3], v[3]) / 2 + 0.6]
            if near:
                v[k] = min(near, key=lambda h: abs(h[2] - v[k]))[2]
    for h in H:
        for k in (0, 1):
            near = [v for v in V if v[1] - 1 <= h[2] <= v[2] + 1 and abs(v[0] - h[k]) <= max(h[3], v[3]) / 2 + 0.6]
            if near:
                h[k] = min(near, key=lambda v: abs(v[0] - h[k]))[0]
    # 線の境目
    XL = cluster([v[0] for v in V] + [h[0] for h in H] + [h[1] for h in H], LINE_TOL)
    YL = cluster([h[2] for h in H] + [v[1] for v in V] + [v[2] for v in V], LINE_TOL)
    items = text_items(chars, H, V, sh.get('bold', []), sh.get('boldFonts', []))
    # spec で足す文字（LibreOffice の描き方が Word と違う所を直す：折り返した題名・ほかのページにある見出しなど）
    for t in sh.get('texts', []):
        x0, top, x1, bottom = t['box']
        items.append({'kind': 'h', 'text': t['text'], 'x0': x0, 'x1': x1, 'top': top, 'bottom': bottom, 'cy': (top + bottom) / 2, 's': t.get('size', 10.5),
                      'f': t.get('font', ''), 'b': t.get('bold', False), 'fixed': True, 'h': t.get('h', 'left'), 'v': t.get('v', 'center'), 'shape': t.get('shape', False)})
    # 線だけのグリッドで囲まれた欄を求める（欄の場所 at で選ぶ・文字のない欄を結合する）
    def edges_on(XB, YB):
        hb, vb = {}, {}
        for x0, x1, y, lw, dash in H:
            j = nearest(YB, y)
            i0, i1 = nearest(XB, x0), nearest(XB, x1)
            for i in range(i0, i1):
                st = side_of(lw, dash)
                if WEIGHT[st] >= WEIGHT.get(hb.get((j, i)), -1):
                    hb[(j, i)] = st
        for x, y0, y1, lw, dash in V:
            i = nearest(XB, x)
            j0, j1 = nearest(YB, y0), nearest(YB, y1)
            for j in range(j0, j1):
                st = side_of(lw, dash)
                if WEIGHT[st] >= WEIGHT.get(vb.get((i, j)), -1):
                    vb[(i, j)] = st
        return hb, vb

    hb0, vb0 = edges_on(XL, YL)
    regs = [(XL[a], YL[b], XL[c], YL[d]) for a, b, c, d in regions(XL, YL, hb0, vb0)]

    if sh.get('analyze'):
        # 作る前に、文字と囲まれた欄の位置だけを返す（欄の場所を見出しの文字から決めるとき）
        return {'items': items, 'regs': regs}
    if sh.get('dump'):
        for r in sorted(regs, key=lambda r: (round(r[1]), r[0])):
            inside = [it['text'] for it in items if r[0] <= it.get('vx', (it['x0'] + it['x1']) / 2) <= r[2] and r[1] <= it['cy'] <= r[3]]
            print(f'  [{r[0]:.0f},{r[1]:.0f},{r[2]:.0f},{r[3]:.0f}] {" / ".join(inside)[:70]}', file=sys.stderr)
        for it in sorted(items, key=lambda it: (round(it['top']), it['x0'])):
            print(f'  ・{it["text"][:40]} ({it["x0"]:.0f},{it["top"]:.0f},{it["x1"]:.0f},{it["bottom"]:.0f}) {it["s"]}', file=sys.stderr)

    def region_at(x, y):
        best = None
        for r in regs:
            if r[0] <= x <= r[2] and r[1] <= y <= r[3]:
                if best is None or (r[2] - r[0]) * (r[3] - r[1]) < (best[2] - best[0]) * (best[3] - best[1]):
                    best = r
        return best

    # 文字の結合セルの範囲（pt）
    def snap_left(x):
        cands = [b for b in XL if 0 <= x - b <= 6]
        return max(cands) if cands else x - 1.0

    def snap_right(x):
        cands = [b for b in XL if 0 <= b - x <= 8]
        return min(cands) if cands else x + 3.0

    boxes = []
    for it in items:
        if it.get('fixed'):
            it['box'] = [it['x0'], it['top'], it['x1'], it['bottom']]
            continue
        if it['kind'] == 'v':
            r = region_at(it['vx'], it['cy'])
            if r and r[2] - r[0] < 4 * it['s']:
                it['box'] = list(r)
            else:
                it['box'] = [snap_left(it['x0']), it['top'] - 1, snap_right(it['x1']), it['bottom'] + 1]
            continue
        s = it['s']
        top, bottom = it['cy'] - 0.62 * s, it['cy'] + 0.62 * s
        # 上下の線が近ければ線まで
        for b in YL:
            if abs(b - top) < 0.3 * s and h_between(H, b, b, (it['x0'] + it['x1']) / 2):
                top = b
            if abs(b - bottom) < 0.3 * s and h_between(H, b, b, (it['x0'] + it['x1']) / 2):
                bottom = b
        # 文字の上下に線があれば、線を越えない（線に寄せた側にそろえ、反対側に字の高さの 1.3 倍までのばす。Excel は字の大きさの約 1.2 倍の高さが要る）
        up = max([h[2] for h in H if h[0] <= it['x1'] and h[1] >= it['x0'] and h[2] < it['cy']], default=-1e9)
        down = min([h[2] for h in H if h[0] <= it['x1'] and h[1] >= it['x0'] and h[2] > it['cy']], default=1e9)
        it['v'] = 'center'
        if bottom > down and top < up:
            top, bottom = up, down
        elif bottom > down:
            bottom = down
            top = max(up, min(top, bottom - 1.3 * s))
            it['v'] = 'bottom'
        elif top < up:
            top = up
            bottom = min(down, max(bottom, top + 1.3 * s))
            it['v'] = 'top'
        it['box'] = [snap_left(it['x0']), top, snap_right(it['x1']), bottom]
    # 上下に重なる横書きのかたまり（行の間がせまい）は、まん中で分ける
    hs = [it for it in items if it['kind'] == 'h']
    for a in hs:
        for b in hs:
            if a is b or a['cy'] >= b['cy']:
                continue
            if a['box'][0] < b['box'][2] and b['box'][0] < a['box'][2] and a['box'][3] > b['box'][1]:
                mid = (a['cy'] + b['cy']) / 2
                a['box'][3] = min(a['box'][3], mid)
                b['box'][1] = max(b['box'][1], mid)
    # 横に重なる（すき間がせまい）ものは、まん中で分ける
    for a in hs:
        for b in hs:
            if a is b or a['x0'] >= b['x0']:
                continue
            if a['box'][1] < b['box'][3] and b['box'][1] < a['box'][3] and a['box'][2] > b['box'][0]:
                mid = (a['x1'] + b['x0']) / 2
                a['box'][2] = min(a['box'][2], mid)
                b['box'][0] = max(b['box'][0], mid)
    # アプリが書く欄
    fields = []
    for f in sh.get('fields', []):
        if 'box' in f:
            box = list(f['box'])
        else:
            r = region_at(*f['at'])
            if not r:
                sys.exit(f'{sh["name"]}: 欄 {f["name"]} の {f["at"]} に囲まれた欄がない')
            box = list(r)
        for k, idx in (('x0', 0), ('top', 1), ('x1', 2), ('bottom', 3)):
            if k in f:
                box[idx] = f[k]
        fields.append((f, box))
    # 文字の箱がアプリの欄にかかるときは、字にかからない所までちぢめる（欄の左右・上下の文字：「電話番号」の右の欄など）
    for it in items:
        if it.get('fixed') or it['kind'] != 'h':
            continue
        b = it['box']
        for _, fb in fields:
            if not (b[0] < fb[2] and fb[0] < b[2] and b[1] < fb[3] and fb[1] < b[3]):
                continue
            if it['x1'] <= fb[0] + 0.3:
                b[2] = min(b[2], fb[0])
            elif it['x0'] >= fb[2] - 0.3:
                b[0] = max(b[0], fb[2])
            elif it['bottom'] <= fb[1] + 0.5:
                b[3] = min(b[3], fb[1])
            elif it['top'] >= fb[3] - 0.5:
                b[1] = max(b[1], fb[3])
    # 境目：線の境目＋文字・欄の境目（線の境目の近くは線に寄せる）
    extra_x = [it['box'][0] for it in items] + [it['box'][2] for it in items] + [b[0] for _, b in fields] + [b[2] for _, b in fields] + [f[0] for f in fills] + [f[2] for f in fills]
    extra_y = [it['box'][1] for it in items] + [it['box'][3] for it in items] + [b[1] for _, b in fields] + [b[3] for _, b in fields] + [f[1] for f in fills] + [f[3] for f in fills]

    def merge_bounds(L, extra):
        add = [e for e in extra if all(abs(e - b) > MIN_GAP for b in L)]
        out = sorted(L + cluster(add, MIN_GAP))
        # 文字の境目どうしが近すぎるとき（線の境目とのあいだ）も詰める
        res = [out[0]]
        for b in out[1:]:
            if b - res[-1] < MIN_GAP:
                if b in L and res[-1] not in L:
                    res[-1] = b
                continue
            res.append(b)
        return res

    XB = merge_bounds(XL, extra_x)
    YB = merge_bounds(YL, extra_y)
    hb, vb = edges_on(XB, YB)
    regs_final = regions(XB, YB, hb, vb)
    # 列の幅・行の高さ（画素の整数に丸める。ＭＳ Ｐゴシック 11pt の数字の幅＝8画素、1画素＝0.75pt）
    px = [round((b - XB[0]) / 0.75) for b in XB]
    for i in range(len(XB) - 1):
        w = max(1, px[i + 1] - px[i])
        ws.column_dimensions[get_column_letter(i + 1)].width = w / 8
    qy = [round((b - YB[0]) / 0.75) for b in YB]
    for j in range(len(YB) - 1):
        ws.row_dimensions[j + 1].height = max(1, qy[j + 1] - qy[j]) * 0.75
    nx, ny = len(XB) - 1, len(YB) - 1
    # 塗り
    for x0, t, x1, b, color in fills:
        for j in range(nearest(YB, t), nearest(YB, b)):
            for i in range(nearest(XB, x0), nearest(XB, x1)):
                ws.cell(j + 1, i + 1).fill = PatternFill('solid', fgColor=color)
    merged = []

    def overlaps(a, b):
        return a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3]

    def merge(i0, j0, i1, j1):
        rng = (i0, j0, i1, j1)
        if i1 - i0 <= 0 or j1 - j0 <= 0:
            return None
        if any(overlaps(rng, m) for m in merged):
            return None
        merged.append(rng)
        if i1 - i0 > 1 or j1 - j0 > 1:
            ws.merge_cells(start_row=j0 + 1, start_column=i0 + 1, end_row=j1, end_column=i1)
        return rng

    def ref(rng):
        i0, j0, i1, j1 = rng
        a = f'{get_column_letter(i0 + 1)}{j0 + 1}'
        b = f'{get_column_letter(i1)}{j1}'
        return a if a == b else f'{a}:{b}'

    fmap = sh.get('fonts', {})
    base = sh.get('font', 'ＭＳ ゴシック')
    out_fields = {}
    for f, box in fields:
        i0, j0, i1, j1 = nearest(XB, box[0]), nearest(YB, box[1]), nearest(XB, box[2]), nearest(YB, box[3])
        if f.get('merge', True) is False:
            # 結合しない範囲（写真・平面図の図を貼る所）：番地だけ
            out_fields[f['name']] = ref((i0, j0, max(i1, i0 + 1), max(j1, j0 + 1)))
            continue
        rng = merge(i0, j0, i1, j1)
        if not rng:
            sys.exit(f'{sh["name"]}: 欄 {f["name"]} がほかの欄と重なる')
        cell = ws.cell(j0 + 1, i0 + 1)
        if 'value' in f:
            cell.value = f['value']
        cell.font = Font(name=f.get('font', base), size=f.get('size', 10.5), family=3, charset=128)
        cell.alignment = Alignment(horizontal=f.get('h', 'left'), vertical=f.get('v', 'center'), wrap_text=f.get('wrap', False),
                                   indent=f.get('indent', 0), shrink_to_fit=f.get('shrink', False))
        out_fields[f['name']] = ref(rng)
    out_items = []
    shapes = []

    def anchor(x, y):
        # pt → 列・行と、その中のずれ（EMU。1画素＝9525）
        X = (x - XB[0]) / 0.75
        Y = (y - YB[0]) / 0.75
        i = max(0, min(len(px) - 2, max(k for k in range(len(px)) if px[k] <= X) if X >= 0 else 0))
        j = max(0, min(len(qy) - 2, max(k for k in range(len(qy)) if qy[k] <= Y) if Y >= 0 else 0))
        return i, round((X - px[i]) * 9525), j, round((Y - qy[j]) * 9525)

    for it in sorted(items, key=lambda it: (it['kind'] != 'v', it['top'], it['x0'])):
        font = fmap.get(it['f'], base) if it['f'] else base
        box = it['box']
        # 線のあいだがせまく、字の高さが入らない（Excel は字の上に1〜2pt あける）→ 図形（テキストボックス）で元の位置に
        tight = it['kind'] == 'h' and (box[3] - box[1]) < it['s'] + 0.5
        rng = None
        if not tight and not it.get('shape'):
            i0, j0 = nearest(XB, box[0]), nearest(YB, box[1])
            i1, j1 = nearest(XB, box[2]), nearest(YB, box[3])
            i1, j1 = max(i1, i0 + 1), max(j1, j0 + 1)
            rng = merge(i0, j0, i1, j1)
        if rng:
            cell = ws.cell(rng[1] + 1, rng[0] + 1)
            cell.value = it['text']
            cell.font = Font(name=font, size=it['s'], bold=it['b'], family=3, charset=128)
            if it['kind'] == 'v':
                spread = (it['bottom'] - it['top']) > len(it['text']) * it['s'] * 1.4
                cell.alignment = Alignment(horizontal='center', vertical='distributed' if spread else 'center', text_rotation=255, wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal=it.get('h', 'left'), vertical=it.get('v', 'center'))
            out_items.append({'text': it['text'], 'ref': ref(rng), 'kind': it['kind'], 'x0': round(it['x0'], 1), 'top': round(it['top'], 1)})
            continue
        # 図形：字の箱（上下は字の高さ＋少し、左右は字の幅＋少し）
        if it.get('fixed'):
            x0, t0, x1, b0 = box
        else:
            x0, x1 = it['x0'] - 0.5, it['x1'] + 2.0
            t0, b0 = it['cy'] - 0.6 * it['s'], it['cy'] + 0.6 * it['s']
        shapes.append({'text': it['text'], 'from': anchor(x0, t0), 'to': anchor(x1, b0), 'w': round((x1 - x0) / 0.75 * 9525), 'h': round((b0 - t0) / 0.75 * 9525),
                       'x': round((x0 - XB[0]) / 0.75 * 9525), 'y': round((t0 - YB[0]) / 0.75 * 9525),
                       'font': font, 'size': it['s'], 'bold': it['b'], 'vertical': it['kind'] == 'v', 'ha': it.get('h', 'left'), 'v': it.get('v', 'center')})
        out_items.append({'text': it['text'], 'ref': 'shape', 'kind': it['kind'], 'x0': round(it['x0'], 1), 'top': round(it['top'], 1)})
    # 文字のない囲まれた欄は結合（アプリが書けるように）
    for i0, j0, i1, j1 in regs_final:
        if (i1 - i0) * (j1 - j0) > 1 and not any(overlaps((i0, j0, i1, j1), m) for m in merged):
            merge(i0, j0, i1, j1)
    # 罫線
    for j in range(ny):
        for i in range(nx):
            t, b, l, r = hb.get((j, i)), hb.get((j + 1, i)), vb.get((i, j)), vb.get((i + 1, j))
            if t or b or l or r:
                ws.cell(j + 1, i + 1).border = Border(top=Side(style=t) if t else Side(), bottom=Side(style=b) if b else Side(),
                                                      left=Side(style=l) if l else Side(), right=Side(style=r) if r else Side())
    # 印刷：A4・余白は原本の位置・1ページに合わせる
    ws.sheet_view.showGridLines = False
    ws.page_setup.paperSize = 9
    ws.page_setup.orientation = 'landscape' if sh.get('landscape') else 'portrait'
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    pw = 841.89 if sh.get('landscape') else 595.3
    ph = 595.3 if sh.get('landscape') else 841.89
    m = ws.page_margins
    m.left = max(0.1, XB[0] / 72)
    m.right = max(0.1, (pw - XB[-1]) / 72)
    m.top = max(0.1, YB[0] / 72)
    m.bottom = max(0.1, (ph - YB[-1]) / 72) if YB[-1] < ph else 0.1
    m.header = m.footer = 0
    ws.print_area = f'A1:{get_column_letter(nx)}{ny}'
    return {'fields': out_fields, 'items': out_items, 'shapes': shapes, 'cols': nx, 'rows': ny, 'xb': [round(b, 2) for b in XB], 'yb': [round(b, 2) for b in YB]}


def esc(t):
    return t.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('"', '&quot;')


def add_shapes(path, shapes_by_sheet):
    """openpyxl は図形を書けないので、保存したファイルに図形（テキストボックス）の部品を足す"""
    import re
    import zipfile
    src = zipfile.ZipFile(path)
    files = {n: src.read(n) for n in src.namelist()}
    src.close()
    wb = files['xl/workbook.xml'].decode()
    rels = files['xl/_rels/workbook.xml.rels'].decode()
    ct = files['[Content_Types].xml'].decode()
    n = 0
    for m in re.finditer(r'<sheet [^>]*name="([^"]*)"[^>]*r:id="([^"]+)"', wb):
        name, rid = m.group(1), m.group(2)
        shapes = shapes_by_sheet.get(name.replace('&amp;', '&'))
        if not shapes:
            continue
        target = re.search(r'Id="%s"[^>]*Target="([^"]+)"' % rid, rels) or re.search(r'Target="([^"]+)"[^>]*Id="%s"' % rid, rels)
        sheet_path = 'xl/' + target.group(1).lstrip('/').replace('xl/', '')
        n += 1
        anchors = []
        for k, sp in enumerate(shapes):
            fc, fco, fr, fro = sp['from']
            tc, tco, tr, tro = sp['to']
            # 字のそろえは 'ha'（'h' は図形の高さ。前は同じ名前にして高さが「left」になり、Excel で開くと修復になる xlsx ができていた。2026-10-29）
            algn = {'left': 'l', 'center': 'ctr', 'right': 'r'}[sp['ha']]
            anc = {'top': 't', 'center': 'ctr', 'bottom': 'b'}[sp['v']]
            vert = ' vert="eaVert"' if sp['vertical'] else ''
            b = ' b="1"' if sp['bold'] else ''
            anchors.append(
                f'<xdr:twoCellAnchor editAs="absolute"><xdr:from><xdr:col>{fc}</xdr:col><xdr:colOff>{fco}</xdr:colOff><xdr:row>{fr}</xdr:row><xdr:rowOff>{fro}</xdr:rowOff></xdr:from>'
                f'<xdr:to><xdr:col>{tc}</xdr:col><xdr:colOff>{tco}</xdr:colOff><xdr:row>{tr}</xdr:row><xdr:rowOff>{tro}</xdr:rowOff></xdr:to>'
                f'<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="{k + 2}" name="テキスト ボックス {k + 1}"/><xdr:cNvSpPr txBox="1"/></xdr:nvSpPr>'
                f'<xdr:spPr><a:xfrm><a:off x="{sp["x"]}" y="{sp["y"]}"/><a:ext cx="{sp["w"]}" cy="{sp["h"]}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></xdr:spPr>'
                f'<xdr:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="{anc}"{vert} rtlCol="0"/><a:lstStyle/>'
                f'<a:p><a:pPr algn="{algn}"/><a:r><a:rPr lang="ja-JP" altLang="en-US" sz="{round(sp["size"] * 100)}"{b}><a:solidFill><a:srgbClr val="000000"/></a:solidFill>'
                f'<a:latin typeface="{esc(sp["font"])}"/><a:ea typeface="{esc(sp["font"])}"/></a:rPr><a:t>{esc(sp["text"])}</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/></xdr:twoCellAnchor>')
        files[f'xl/drawings/drawing{n}.xml'] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            + ''.join(anchors) + '</xdr:wsDr>').encode()
        ct = ct.replace('</Types>', f'<Override PartName="/xl/drawings/drawing{n}.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>')
        rel_path = sheet_path.replace('worksheets/', 'worksheets/_rels/') + '.rels'
        rel = f'<Relationship Id="rIdShapes" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing{n}.xml"/>'
        if rel_path in files:
            files[rel_path] = files[rel_path].decode().replace('</Relationships>', rel + '</Relationships>').encode()
        else:
            files[rel_path] = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">' + rel + '</Relationships>').encode()
        xml = files[sheet_path].decode()
        if 'xmlns:r=' not in xml[:600]:
            xml = xml.replace('<worksheet ', '<worksheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ', 1)
        # drawing は pageSetup・headerFooter・rowBreaks・colBreaks のあと
        pos = max(xml.rfind(t) for t in ('</headerFooter>', '<pageSetup', '</rowBreaks>', '</colBreaks>', '<pageMargins'))
        end = xml.index('>', pos) + 1
        if xml[pos:].startswith('<pageSetup') or xml[pos:].startswith('<pageMargins'):
            end = xml.index('/>', pos) + 2
        xml = xml[:end] + '<drawing r:id="rIdShapes"/>' + xml[end:]
        files[sheet_path] = xml.encode()
    files['[Content_Types].xml'] = ct.encode()
    out = zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED)
    for name in ['[Content_Types].xml'] + [k for k in files if k != '[Content_Types].xml']:
        out.writestr(name, files[name])
    out.close()


def analyze(sh, pdf_path=None):
    """欄の場所を決めるための、文字（text・x0・x1・top・bottom・s）と囲まれた欄（x0, top, x1, bottom）"""
    return build_sheet(None, pdfplumber.open(pdf_path or sh['pdf']), {**sh, 'analyze': True})


def main():
    spec = json.load(open(sys.argv[1], encoding='utf-8'))
    wb = Workbook()
    wb.remove(wb.active)
    # 標準の字体（列の幅の数字の幅 8 画素の基準）
    normal = wb._named_styles['Normal']
    normal.font = Font(name='ＭＳ Ｐゴシック', size=11)
    wb._fonts[0] = Font(name='ＭＳ Ｐゴシック', size=11)
    result = {}
    for sh in spec['sheets']:
        pdf = pdfplumber.open(sh.get('pdf', spec.get('pdf')))
        ws = wb.create_sheet(sh['name'])
        result[sh['name']] = build_sheet(ws, pdf, sh)
        print(f'{sh["name"]}: {result[sh["name"]]["cols"]}列 × {result[sh["name"]]["rows"]}行・文字 {len(result[sh["name"]]["items"])}', file=sys.stderr)
    wb.save(spec['out'])
    add_shapes(spec['out'], {name: r['shapes'] for name, r in result.items()})
    json.dump(result, open(spec['map'], 'w', encoding='utf-8'), ensure_ascii=False, indent=1)


if __name__ == '__main__':
    main()
